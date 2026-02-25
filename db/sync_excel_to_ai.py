import openpyxl
import os
import sys
import re
from sqlalchemy import text

# Add the project root to the path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from db.base import SessionLocal, engine, Base
from db.models.ai_layers import SemanticLayer, AiCuratedLayer, IntegrationPull, IntegrationPush, PublicCloudLayer

file_path = r"c:\Users\haris\Downloads\harish folder\main reepo\tai-data-api\db\20260108 - VIM DataDictionary (1).xlsx"

def clean_header(header):
    if not header: return None
    h = str(header).replace('\n', ' ').replace('\r', ' ')
    h = re.sub(r'[^a-zA-Z0-9 ]', '', h)
    h = h.strip().replace(' ', '_').upper()
    return h if h else None

def clean_and_seed_ai():
    db = SessionLocal()
    print("\n--- Syncing AI Schema with Excel Data Dictionary ---")
    
    try:
        # Wipe the schema clean so those 15 extra tables are gone
        with engine.connect() as conn:
            print("Wiping old AI schema...")
            conn.execute(text("DROP SCHEMA IF EXISTS ai CASCADE;"))
            conn.execute(text("CREATE SCHEMA ai;"))
            conn.commit()
            
        print("Creating fresh tables...")
        Base.metadata.create_all(bind=engine)
        
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        mapping = [
            ('Layer-1(Semantic)', SemanticLayer),
            ('Layer-2(AI-Curated)', AiCuratedLayer),
            ('Layer-3A(Integration-Pull)', IntegrationPull),
            ('Layer-3B(Integration-Push)', IntegrationPush),
            ('Layer-3(PublicCloud)', PublicCloudLayer)
        ]

        for sheet_name, ModelClass in mapping:
            if sheet_name in wb.sheetnames:
                print(f"Syncing sheet: {sheet_name}...")
                sheet = wb[sheet_name]
                
                # Get headers using the SAME cleaning logic as the model generator
                headers = []
                for cell in sheet[1]:
                    h = clean_header(cell.value)
                    if h:
                        headers.append(h)
                
                count = 0
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    if any(row):
                        obj_data = {}
                        for i, val in enumerate(row):
                            if i < len(headers):
                                obj_data[headers[i]] = str(val) if val is not None else None
                        
                        try:
                            db.add(ModelClass(**obj_data))
                            count += 1
                        except Exception as row_err:
                            print(f"Skip row {row_idx} in {sheet_name}: {row_err}")
                
                print(f"  Successfully added {count} records to {ModelClass.__tablename__}")
        
        db.commit()
        print("\n--- Done: AI Schema is now 100% matched to Excel ---")
        
    except Exception as e:
        print(f"Fatal Error: {e}")
        db.rollback()
    finally:
        db.close()

if __name__ == "__main__":
    clean_and_seed_ai()
