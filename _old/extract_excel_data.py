import openpyxl
import os
import json

file_path = r"c:\Users\haris\Downloads\harish folder\main reepo\tai-data-api\db\20260108 - VIM DataDictionary (1).xlsx"

def extract_data():
    if not os.path.exists(file_path):
        print("Excel not found")
        return

    wb = openpyxl.load_workbook(file_path, data_only=True)
    dict_data = {}
    
    sheets_to_process = {
        'Layer-1(Semantic)': 'semantic_layer',
        'Layer-2(AI-Curated)': 'ai_curated_layer',
        'Layer-3A(Integration-Pull)': 'integration_pull',
        'Layer-3B(Integration-Push)': 'integration_push',
        'Layer-3(PublicCloud)': 'public_cloud_layer'
    }

    for sheet_name, table_name in sheets_to_process.items():
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            # Get headers from first row
            headers = [cell.value for cell in sheet[1] if cell.value]
            
            rows = []
            # Start from row 2
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(row): # only if row is not empty
                    data_row = {}
                    for i, val in enumerate(row):
                        if i < len(headers):
                            data_row[headers[i]] = val
                    rows.append(data_row)
            
            dict_data[table_name] = {
                "headers": headers,
                "data": rows
            }
    
    print(json.dumps(dict_data, indent=2))

if __name__ == "__main__":
    extract_data()
